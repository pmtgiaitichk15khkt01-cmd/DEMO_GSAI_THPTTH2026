# VIP 76.0 - HỆ THỐNG GIA SƯ AI THPT 2026 (CHUẨN GDPT 2018 - EDTECH & PSYCHOMETRICS SOTA)
# -*- coding: utf-8 -*-
"""
AZOTA & LMS NATIONAL EXPORT SERVICE (XUẤT BẢN ĐỀ CHUẨN AZOTA & LMS QUỐC GIA)
Chuyển đổi 1-chạm toàn bộ đề thi sang định dạng chuẩn nhận diện tự động của Azota / Shub / K12Online:
- Bọc toàn bộ công thức toán học inline $...$ thành $$...$$ (chuẩn LaTeX parser Azota).
- Chuẩn hóa cấu trúc 3 Phần chuẩn Bộ GD&ĐT (Quyết định số 764/QĐ-BGDĐT).
- Tự động sinh Khối Bảng Đáp Án tự động nhận diện ở cuối đề:
    ĐÁP ÁN:
    1.A
    2.B
    ...
- Hỗ trợ xuất cả định dạng Văn bản thuần (.txt) và Microsoft Word (.docx).
"""

import re
import io
from typing import Dict, Any, List, Optional


def convert_math_to_double_dollars(text: str) -> str:
    """
    Chuyển đổi công thức toán học $...$ thành $$...$$:
    - Nếu đã là $$...$$ thì giữ nguyên.
    - Nếu là $...$ thì bọc thành $$...$$.
    """
    if not text:
        return ""
    s = str(text)

    # Bảo vệ các khối $$ đã có bằng token tạm
    protected_blocks = []
    def save_double(m):
        protected_blocks.append(m.group(0))
        return f"__DOUBLE_MATH_TOKEN_{len(protected_blocks)-1}__"

    s = re.sub(r'\$\$(.*?)\$\$', save_double, s, flags=re.DOTALL)

    # Chuyển $inline$ thành $$inline$$
    def wrap_single(m):
        inner = m.group(1).strip()
        return f"$${inner}$$"

    s = re.sub(r'\$([^$\n]+)\$', wrap_single, s)

    # Khôi phục các khối đã có
    for idx, token_content in enumerate(protected_blocks):
        s = s.replace(f"__DOUBLE_MATH_TOKEN_{idx}__", token_content)

    return s


def export_exam_to_azota_text(exam: Dict[str, Any]) -> str:
    """
    Xuất đề thi hoàn chỉnh ra file văn bản chuẩn Azota (.txt).
    """
    if not exam:
        return "Không có dữ liệu đề thi."

    title = exam.get("title") or f"ĐỀ THI {exam.get('subject', 'TOÁN HỌC').upper()} - CHUẨN BỘ GD&ĐT"
    school = exam.get("school", "TRƯỜNG THPT TÂN HIỆP & TRUNG TÂM THIỆN NHÂN")
    year = exam.get("year", "2025 - 2026")
    subject = exam.get("subject", "Toán học")

    lines = []
    lines.append(f"{'='*60}")
    lines.append(f"{school.upper()}")
    lines.append(f"{title.upper()} ({year})")
    lines.append(f"Môn thi: {subject} | Thời gian làm bài: 90 phút (Không kể thời gian phát đề)")
    lines.append(f"{'='*60}\n")

    p1 = exam.get("p1", [])
    p2 = exam.get("p2", [])
    p3 = exam.get("p3", [])
    doc_hieu = exam.get("part_doc_hieu")
    part_viet = exam.get("part_viet")

    # Xử lý đề Ngữ văn
    if doc_hieu or part_viet:
        lines.append("PHẦN I: ĐỌC HIỂU (4.0 điểm)\n")
        if doc_hieu:
            lines.append(convert_math_to_double_dollars(doc_hieu.get("text", "")))
            lines.append("")
            for idx, q in enumerate(doc_hieu.get("questions", [])):
                q_text = convert_math_to_double_dollars(q.get("q", ""))
                lines.append(f"{q_text}\n")

        if part_viet:
            lines.append("\nPHẦN II: VIẾT (6.0 điểm)\n")
            for idx, q in enumerate(part_viet):
                q_text = convert_math_to_double_dollars(q.get("q", ""))
                lines.append(f"{q_text}\n")

        lines.append("\n--- HẾT ---")
        return "\n".join(lines)

    # PHẦN I: TRẮC NGHIỆM 4 LỰA CHỌN
    p1_answers = []
    if p1:
        lines.append("PHẦN I. Thí sinh trả lời từ câu 1 đến câu " + str(len(p1)) + ". Mỗi câu hỏi thí sinh chỉ chọn một phương án.\n")
        for idx, q in enumerate(p1):
            q_num = idx + 1
            raw_q = convert_math_to_double_dollars(q.get("q", ""))
            lines.append(f"Câu {q_num}. {raw_q}")
            
            opts = q.get("opt", [])
            for opt_idx, opt_text in enumerate(opts):
                clean_opt = convert_math_to_double_dollars(opt_text)
                letter = chr(65 + opt_idx)
                # Đảm bảo có tiền tố A. B. C. D.
                if not re.match(r'^[A-D]\.', clean_opt.strip()):
                    clean_opt = f"{letter}. {clean_opt}"
                lines.append(clean_opt)
            lines.append("")
            p1_answers.append(f"{q_num}.{q.get('ans', 'A')}")

    # PHẦN II: TRẮC NGHIỆM ĐÚNG/SAI
    p2_answers = []
    if p2:
        lines.append("\nPHẦN II. Thí sinh trả lời từ câu 1 đến câu " + str(len(p2)) + ". Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn Đúng hoặc Sai.\n")
        for idx, q in enumerate(p2):
            q_num = idx + 1
            raw_q = convert_math_to_double_dollars(q.get("q", ""))
            lines.append(f"Câu {q_num}. {raw_q}")
            
            stmts = q.get("stmts", [])
            stmt_ans_list = []
            stmt_letters = ['a', 'b', 'c', 'd']
            for s_idx, stmt in enumerate(stmts):
                prefix = stmt_letters[s_idx] if s_idx < 4 else f"y{s_idx+1}"
                t_text = convert_math_to_double_dollars(stmt.get("t", ""))
                # Đảm bảo có tiền tố a) b) c) d)
                if not re.match(r'^[a-d]\)', t_text.strip()):
                    t_text = f"{prefix}) {t_text}"
                lines.append(t_text)

                expected_bool = stmt.get("a") if "a" in stmt else stmt.get("ans", True)
                ans_char = "Đ" if expected_bool else "S"
                stmt_ans_list.append(f"{prefix}) {ans_char}")

            lines.append("")
            p2_answers.append(f"Câu {q_num}: " + "  ".join(stmt_ans_list))

    # PHẦN III: TRẢ LỜI NGẮN
    p3_answers = []
    if p3:
        lines.append("\nPHẦN III. Thí sinh trả lời từ câu 1 đến câu " + str(len(p3)) + ". Điền đáp án số vào phiếu trả lời.\n")
        for idx, q in enumerate(p3):
            q_num = idx + 1
            raw_q = convert_math_to_double_dollars(q.get("q", ""))
            lines.append(f"Câu {q_num}. {raw_q}\n")
            p3_answers.append(f"{q_num}. {q.get('ans', '0')}")

    lines.append("\n--- HẾT ---\n")

    # BẢNG ĐÁP ÁN CHUẨN NHẬN DIỆN TỰ ĐỘNG AZOTA
    lines.append(f"{'='*30} ĐÁP ÁN {'='*30}")
    if p1_answers:
        lines.append("PHẦN I:")
        # In hàng ngang 4 đáp án 1 dòng cho gọn
        for i in range(0, len(p1_answers), 4):
            lines.append("  ".join(p1_answers[i:i+4]))
        lines.append("")

    if p2_answers:
        lines.append("PHẦN II:")
        for ans_line in p2_answers:
            lines.append(ans_line)
        lines.append("")

    if p3_answers:
        lines.append("PHẦN III:")
        for i in range(0, len(p3_answers), 3):
            lines.append("  ".join(p3_answers[i:i+3]))
        lines.append("")

    return "\n".join(lines)


def export_exam_to_azota_docx(exam: Dict[str, Any]) -> io.BytesIO:
    """
    Xuất đề thi ra tệp Microsoft Word (.docx) chuẩn format Azota.
    """
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        # Fallback nếu thiếu thư viện docx
        bio = io.BytesIO()
        text_content = export_exam_to_azota_text(exam)
        bio.write(text_content.encode('utf-8'))
        bio.seek(0)
        return bio

    doc = Document()

    # Thiết lập lề A4
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Tiêu đề
    title = exam.get("title") or f"ĐỀ THI {exam.get('subject', 'TOÁN HỌC').upper()} - CHUẨN BỘ GD&ĐT"
    school = exam.get("school", "TRƯỜNG THPT TÂN HIỆP & TRUNG TÂM THIỆN NHÂN")
    
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sch = p_head.add_run(f"{school.upper()}\n")
    run_sch.bold = True
    run_sch.font.size = Pt(12)

    run_title = p_head.add_run(f"{title.upper()}\n")
    run_title.bold = True
    run_title.font.size = Pt(14)
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_info.add_run(f"Môn: {exam.get('subject', 'Toán học')}  |  Thời gian làm bài: 90 phút (Chuẩn QĐ 764/QĐ-BGDĐT)")
    run_sub.font.italic = True
    run_sub.font.size = Pt(10.5)

    doc.add_paragraph("─" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Thêm nội dung text
    text_content = export_exam_to_azota_text(exam)
    # Tách đoạn và thêm vào docx
    for line in text_content.split("\n"):
        line_s = line.strip()
        if not line_s:
            continue
        p = doc.add_paragraph()
        if line_s.startswith("PHẦN ") or "ĐÁP ÁN" in line_s:
            run = p.add_run(line_s)
            run.bold = True
            run.font.size = Pt(11.5)
            if "ĐÁP ÁN" in line_s:
                run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        elif line_s.startswith("Câu ") and ":" not in line_s[:10]:
            run = p.add_run(line_s)
            run.bold = True
            run.font.size = Pt(11)
        else:
            run = p.add_run(line_s)
            run.font.size = Pt(11)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
