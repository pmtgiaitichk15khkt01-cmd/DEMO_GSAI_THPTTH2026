"""Regression cases reproduced during the live four-station audit."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
from qa_runtime import (admin_password_matches, split_api_keys, numeric_answer_matches,
                        normalize_bool, paired_statistics, display_frame, parse_grade,
                        parse_math_expression, exam_response_metadata, observed_item_matrix)


class TestValidatedBoundaries(unittest.TestCase):
    def test_keys_are_separate_and_deduplicated(self):
        self.assertEqual(split_api_keys("one, two;one\nthree"), ["one", "two", "three"])

    def test_admin_requires_configured_secret(self):
        self.assertFalse(admin_password_matches("GiaoVienKHKT@2026", ""))
        self.assertFalse(admin_password_matches("GiaoVienKHKT@2026", "private-test-password"))
        self.assertTrue(admin_password_matches("private-test-password", "private-test-password"))

    def test_integer_answers_do_not_enable_implicit_rounding(self):
        for u, e in [("3.4", "3"), ("2.4", "2"), ("100009", "100000")]:
            with self.subTest(user=u): self.assertFalse(numeric_answer_matches(u, e))

    def test_zero_and_explicit_rounding(self):
        self.assertTrue(numeric_answer_matches(0, 0))
        self.assertTrue(numeric_answer_matches("1,25", "1.25"))
        self.assertTrue(numeric_answer_matches("1.25", "1.3", 1))
        for value in (None, "", "NaN", "inf", "A"):
            with self.subTest(value=value): self.assertFalse(numeric_answer_matches(value, "0"))

    def test_tolerance_boundary(self):
        self.assertTrue(numeric_answer_matches("2.015", "2"))
        self.assertFalse(numeric_answer_matches("2.0151", "2"))

    def test_mcq_without_letter_prefixes(self):
        from qa_runtime import mcq_answer_matches
        q = {"opt": ["$x=1$", "$x=2$", "$x=3$", "$x=4$"], "ans": "B"}
        self.assertTrue(mcq_answer_matches(q, "$x=2$"))
        self.assertFalse(mcq_answer_matches(q, "$x=1$"))

    def test_demo_hash_is_overridden_by_configured_password(self):
        import hashlib
        configured = "sha256:" + hashlib.sha256(b"test-only").hexdigest()
        self.assertTrue(admin_password_matches("test-only", configured))
        self.assertFalse(admin_password_matches("test-only", "new-private-password"))

    def test_string_false_is_not_truthy(self):
        from services.scoring_service import grade_part2
        qs = [{"stmts": [{"t": "x", "a": "false"}]}]
        self.assertEqual(grade_part2(qs, {"p2_0_0": "Đúng"})[0], 0)
        self.assertEqual(grade_part2(qs, {"p2_0_0": "Sai"})[0], .1)
        self.assertIsNone(normalize_bool("unknown"))

    def test_grade_parsing(self):
        for value in (10, "10", "Lớp 10"):
            self.assertEqual(parse_grade(value), 10)
        for value in (5, 13, None, "10.5"):
            with self.assertRaises(ValueError): parse_grade(value)

    def test_constant_gain_has_stable_result_schema(self):
        for pre, post in [([4, 5, 6], [5, 6, 7]), ([4, 5], [4, 5])]:
            result = paired_statistics(pre, post)
            self.assertEqual(set(result), {"n", "gain", "hake", "t", "p", "dz", "ci", "shapiro_p", "wilcoxon_p"})
            self.assertIsNone(result["p"])

    def test_mixed_log_objects_are_arrow_serializable(self):
        import pyarrow as pa
        frame = pd.DataFrame({"detail": [5, "Tính sai", {"a": 1}, [1, 2], None]})
        rendered = display_frame(frame)
        pa.Table.from_pandas(rendered)
        self.assertEqual(frame.iloc[0, 0], 5)
        self.assertTrue(all(isinstance(v, str) for v in rendered["detail"]))

    def test_math_parser_blocks_code_and_expensive_syntax(self):
        for expr in ["__import__('os')", "sp.__dict__", "x.__class__", "[x for x in (1,2)]", "2**100000", "x**(2**10)", "x" * 301]:
            with self.subTest(expr=expr): self.assertIsNone(parse_math_expression(expr))
        self.assertEqual(str(parse_math_expression("x**3 - 3*x")), "x**3 - 3*x")

    def test_stationary_point_is_not_always_extremum(self):
        from services.neuro_symbolic_verifier import verify_calculus_function
        cubic = verify_calculus_function("x^3")
        self.assertTrue(cubic["verified"])
        self.assertEqual(cubic["stationary_points"], [0.0])
        self.assertEqual(cubic["critical_points"], [])
        self.assertEqual(verify_calculus_function("x^4")["critical_points"], [0.0])
        self.assertFalse(verify_calculus_function("sin(x)")["verified"])

    def test_wrong_answer_never_gets_verified_badge(self):
        from services.neuro_symbolic_verifier import audit_ai_generated_question_sympy
        question = {"q": "Hàm số $y=x^3$ có bao nhiêu cực trị?", "opt": ["A. 0", "B. 1", "C. 2", "D. 3"], "ans": "B"}
        result = audit_ai_generated_question_sympy(question)
        self.assertFalse(result["verified"])
        self.assertTrue(result["hallucination_detected"])
        question["ans"] = "A"
        self.assertTrue(audit_ai_generated_question_sympy(question)["verified"])

    def test_integral_verification_and_limits(self):
        from services.neuro_symbolic_verifier import verify_definite_integral
        self.assertEqual(verify_definite_integral("x^2", 0, 3)["numeric_result"], 9)
        self.assertFalse(verify_definite_integral("sin(x)", 0, 1)["verified"])

    def test_empty_review_history_produces_no_reminders(self):
        from services.fsrs_spaced_repetition_service import analyze_student_fsrs_review_schedule
        self.assertEqual(analyze_student_fsrs_review_schedule(), [])
        self.assertEqual(analyze_student_fsrs_review_schedule([]), [])

    def test_session_defaults_are_not_shared(self):
        from services.state_service import init_centralized_session_state
        class Session:
            def __init__(self): self.session_state = {}
        first, second = Session(), Session()
        init_centralized_session_state(first)
        init_centralized_session_state(second)
        first.session_state["exam_answers"]["x"] = 1
        self.assertEqual(second.session_state["exam_answers"], {})

    def test_observed_matrix_uses_correctness_and_exam_version(self):
        exam = {"p1": [{"q": "one", "ans": "A"}], "p3": [{"q": "two", "ans": "3"}]}
        logs = []
        for i, ans in enumerate([{"p1_0": "A", "p3_0": "3.4"}, {"p1_0": "B", "p3_0": "3"}, {"p1_0": "A", "p3_0": "3"}]):
            logs.append({"type": "EXAM_RESULT", "event_id": str(i), "answers": {**ans, **exam_response_metadata(exam, ans)}})
        matrix, names = observed_item_matrix(logs + [logs[0]])
        self.assertEqual(matrix.tolist(), [[1, 0], [0, 1], [1, 1]])
        self.assertEqual(names, ["p1_0", "p3_0"])
        self.assertEqual(observed_item_matrix([{"answers": {"p1_0": "A"}}])[0].size, 0)

    def test_negative_alpha_is_preserved(self):
        from services.psychometrics_service import compute_cronbach_alpha
        matrix = np.array([[1, 0], [0, 1], [1, 0], [0, 0]], dtype=float)
        self.assertLess(compute_cronbach_alpha(matrix)["alpha"], 0)


class TestDatabaseEvents(unittest.TestCase):
    def setUp(self):
        from database import DatabaseManager
        self.temp = tempfile.TemporaryDirectory()
        self.db = DatabaseManager(Path(self.temp.name) / "qa.db")

    def tearDown(self):
        self.db._get_connection().close()
        self.temp.cleanup()

    def test_missing_event_ids_never_overwrite_each_other(self):
        first = {"name": "QA", "grade": 10, "score": 3}
        second = {"name": "QA", "grade": "10", "score": 4}
        self.assertTrue(self.db.save_exam_attempt(first))
        self.assertTrue(self.db.save_exam_attempt(second))
        rows = self.db._get_connection().execute("SELECT id,grade_level,student_id FROM exam_attempts").fetchall()
        self.assertEqual(len(rows), 2)
        self.assertEqual([r["grade_level"] for r in rows], [10, 10])
        self.assertNotEqual(rows[0]["student_id"], rows[1]["student_id"])

    def test_replay_does_not_change_original_score(self):
        event = {"event_id": "repeat", "student_id": "qa", "grade": 12, "score": 3}
        self.db.save_exam_attempt(event)
        event["score"] = 9
        self.db.save_exam_attempt(event)
        self.assertEqual(self.db._get_connection().execute("SELECT score FROM exam_attempts").fetchone()[0], 3)

    def test_invalid_score_and_grade_are_rejected(self):
        self.assertFalse(self.db.save_exam_attempt({"score": 11, "grade": 12}))
        self.assertFalse(self.db.save_exam_attempt({"score": 3, "grade": 13}))

    def test_cloud_parent_precedes_child(self):
        with patch.object(self.db, "_sync_to_supabase", return_value=True) as sync:
            self.assertTrue(self.db._sync_parent_and_event("exam_attempts", {"student_id": "qa", "grade_level": 12}, "url", "key"))
            self.assertEqual([c.args[0] for c in sync.call_args_list], ["students", "exam_attempts"])

    def test_ambiguous_pairs_are_not_averaged_into_research(self):
        for event_id, phase in [("a", "pre"), ("b", "post"), ("c", "post")]:
            self.db.save_exam_attempt({"event_id": event_id, "student_id": "qa", "grade": 12, "score": 5, "phase": phase, "subject": "Toán học"})
        self.assertEqual(self.db.get_paired_exam_data("Toán học", 12), [])


class TestImportExport(unittest.TestCase):
    def test_word_export_is_a4_and_preserves_answers(self):
        from docx import Document
        from docx.shared import Mm
        from services.azota_lms_exporter import export_exam_to_azota_docx, export_exam_to_azota_text
        exam = {"subject": "Toán học", "duration": 15,
                "p1": [{"q": "Đạo hàm của $x^3$?", "opt": ["A. $3x^2$", "B. $x$", "C. $2x$", "D. $1$"], "ans": "A"}],
                "p2": [{"q": "Xét mệnh đề", "stmts": [{"t": "x", "a": "false"}]}]}
        doc = Document(export_exam_to_azota_docx(exam))
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, places=1)
        self.assertAlmostEqual(doc.sections[0].page_height.mm, 297, places=1)
        text = export_exam_to_azota_text(exam)
        self.assertIn("15 phút", text)
        self.assertIn("a) S", text)
        self.assertIn("$$x^3$$", text)

    def test_docx_import_detects_two_questions(self):
        import io
        from docx import Document
        from exam_file_extractor import extract_exam_file, split_exam_text
        doc = Document()
        for line in ["Câu 1. Đạo hàm của x^3?", "A. 3x^2", "B. x", "C. 2x", "D. 1", "Câu 2. Hàm x^3 có mấy cực trị?", "A. 0", "B. 1", "C. 2", "D. 3"]:
            doc.add_paragraph(line)
        stream = io.BytesIO()
        doc.save(stream)
        extracted = extract_exam_file(stream.getvalue(), "qa.docx")
        self.assertFalse(extracted["needs_ocr"])
        self.assertEqual(len(split_exam_text(extracted["text"])), 2)

    def test_text_pdf_and_scanned_pdf_are_distinguished(self):
        import fitz
        from exam_file_extractor import extract_exam_file
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 50), "Cau 1. Derivative of x^3? A. 3x^2 B. x C. 2x D. 1")
        result = extract_exam_file(doc.tobytes(), "qa.pdf")
        self.assertIn("Derivative", result["text"])
        self.assertFalse(result["needs_ocr"])
        doc.close()


if __name__ == "__main__":
    unittest.main()
